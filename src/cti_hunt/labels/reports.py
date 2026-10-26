"""Đọc file gán nhãn report (bộ nhãn tham chiếu: hiện là gan_nhan_report_silver_v1.xlsx) và kiểm tra HÌNH THỨC.

Chỉ kiểm tra hình thức: mã có trong v19.2 không, bị thu hồi không, tên đúng bản v19.2 chưa, các cột
bắt buộc có giá trị chưa. Không đề xuất mã, không đánh giá nhãn đúng/sai về nội dung.

Cột `report` trong Nhan/IOC là công thức Excel; openpyxl chỉ đọc được giá trị nếu file đã được Excel
lưu lại. Vì vậy mã report luôn lấy từ 2 ký tự đầu của ma_doan.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import warnings

from openpyxl import load_workbook

from ..attack.catalog import AttackCatalog, tactic_key

EXAMPLE_PREFIX = "VD"

# Cột bối cảnh của sheet Nhan (R6). Thứ tự này cũng là thứ tự in ở score-ttp.
CONTEXT_COLUMNS = ("tool", "execution_path", "source_role", "target_role", "account_type",
                   "preceded_by", "command_example")


@dataclass
class Segment:
    ma_doan: str
    report: str
    loai: str          # đoạn văn | dòng bảng | ảnh | tiêu đề | mã/truy vấn
    van_ban: str
    da_xem: bool = False
    anh_hanh_vi: str = ""


def _rows(ws) -> list[dict]:
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    out = []
    for i, r in enumerate(it, start=2):
        d = {h: (v.strip() if isinstance(v, str) else v) for h, v in zip(header, r) if h}
        if not d.get("ma_doan"):
            continue
        d["_row"] = i
        d["report"] = str(d["ma_doan"])[:2]
        out.append(d)
    return out


def load(xlsx: Path) -> dict[str, list]:
    with warnings.catch_warnings():
        warnings.filterwarnings("ignore", message="Data Validation extension")
        wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        segs = [Segment(str(r["ma_doan"]), r["report"], str(r.get("loai") or ""), str(r.get("van_ban") or ""),
                        bool(r.get("da_xem")), str(r.get("anh_hanh_vi") or ""))
                for r in _rows(wb["DoanVan"])]
        nhan = [r for r in _rows(wb["Nhan"]) if r["report"] != EXAMPLE_PREFIX]
        ioc = [r for r in _rows(wb["IOC"]) if r["report"] != EXAMPLE_PREFIX]
    finally:
        wb.close()
    return {"segments": segs, "nhan": nhan, "ioc": ioc}


def gold_pairs(nhan: list[dict], in_scope_only: bool = False) -> set[tuple[str, str]]:
    """Khóa so khớp khi đo khâu ②: (report, technique_id) (R9, R13)."""
    return {(r["report"], str(r["technique_id"]).strip().upper()) for r in nhan
            if r.get("technique_id") and (not in_scope_only or r.get("in_scope") == "co")}


def gold_context(nhan: list[dict]) -> dict[tuple[str, str], dict]:
    """Khóa (ma_doan, technique_id) -> dòng nhãn, để đo P/R từng trường context (2b).

    Khóa theo ĐOẠN chứ không theo report như gold_pairs: R1 định nghĩa một dòng = một cặp (đoạn, kỹ thuật),
    nên đây là khóa 1-1 với dòng nhãn. `in_scope` đi kèm vì theo R6 chỉ dòng in_scope=co mới điền context.
    """
    out: dict[tuple[str, str], dict] = {}
    for r in nhan:
        tid = str(r.get("technique_id") or "").strip().upper()
        if not tid:
            continue
        key = (str(r["ma_doan"]).strip(), tid)
        if key in out:                      # trùng cặp (đoạn, kỹ thuật) — check() đã báo lỗi R1
            continue
        out[key] = {"report": r["report"], "in_scope": r.get("in_scope"),
                    **{c: r.get(c) for c in CONTEXT_COLUMNS}}
    return out


def gold_iocs(ioc_rows: list[dict]) -> list[tuple[str, str, str]]:
    """(report, gia_tri, loai) — khóa so khớp khi đo nhánh IOC (R10)."""
    return [(r["report"], str(r["gia_tri"]), str(r.get("loai") or "")) for r in ioc_rows if r.get("gia_tri")]


def accepted_names(catalog: AttackCatalog, tid: str) -> list[str]:
    """Tên hợp lệ theo bản ATT&CK đang dùng. Kỹ thuật con chấp nhận cả tên ngắn ('PowerShell') lẫn dạng
    'Cha: Con' ('Command and Scripting Interpreter: PowerShell') — dạng sau là dạng file gán nhãn dùng (VD-001)
    và là dạng attack.mitre.org hiển thị. Phần tử cuối là dạng đầy đủ, dùng trong thông báo."""
    name = catalog.resolve(tid).name
    if "." not in tid:
        return [name]
    parent = catalog.resolve(tid.split(".")[0])
    return [name, f"{parent.name}: {name}"] if parent.status == "ok" and parent.name else [name]


def check(data: dict, catalog: AttackCatalog) -> list[dict]:
    """Trả danh sách vấn đề hình thức: {muc: loi|canh_bao, sheet, dong, ma_doan, noi_dung}."""
    issues = []

    def add(level, sheet, r, msg):
        issues.append({"muc": level, "sheet": sheet, "dong": r.get("_row"), "ma_doan": r.get("ma_doan"),
                       "noi_dung": msg})

    seg_ids = {s.ma_doan for s in data["segments"]}
    seg_type = {s.ma_doan: s.loai for s in data["segments"]}
    seen = set()
    for r in data["nhan"]:
        tid = str(r.get("technique_id") or "").strip().upper()
        if r["ma_doan"] not in seg_ids:
            add("loi", "Nhan", r, "ma_doan không có trong DoanVan")
        if seg_type.get(r["ma_doan"]) == "tiêu đề":
            add("canh_bao", "Nhan", r, "đoạn loại 'tiêu đề' không gán (R9)")
        if not tid:
            add("loi", "Nhan", r, "thiếu technique_id")
            continue
        key = (r["ma_doan"], tid)
        if key in seen:
            add("loi", "Nhan", r, f"trùng cặp (đoạn, kỹ thuật) {tid} (R1)")
        seen.add(key)
        res = catalog.resolve(tid)
        if res.status == "unknown":
            add("loi", "Nhan", r, f"{tid} không có trong ATT&CK v{catalog.version}")
        elif res.status == "revoked":
            add("loi", "Nhan", r, f"{tid} đã bị thu hồi ở v{catalog.version}; mã thay thế: {res.final} (R14)")
        elif res.status == "deprecated":
            add("loi", "Nhan", r, f"{tid} đã deprecated ở v{catalog.version}")
        else:
            name = str(r.get("technique_name") or "").strip()
            accepted = accepted_names(catalog, tid)
            if name and name not in accepted:
                add("canh_bao", "Nhan", r, f"technique_name khác tên v{catalog.version}: '{accepted[-1]}' (R14)")
            tac = str(r.get("tactic") or "").strip()
            if not tac:
                add("loi", "Nhan", r, "thiếu tactic (R13)")
            elif tactic_key(tac) not in catalog.tactics:
                add("loi", "Nhan", r, f"tactic '{tac}' không có trong danh sách tactic v{catalog.version}")
            elif tactic_key(tac) not in catalog.techniques[tid].tactics and \
                    "tactic theo mục đích" not in str(r.get("ghi_chu") or ""):
                add("canh_bao", "Nhan", r, f"tactic '{tac}' không thuộc tactic chính thức của {tid}; "
                                           "nếu cố ý thì ghi_chu cần 'tactic theo mục đích' (R13)")
        if r.get("muc_ro") not in ("ro", "mo_ho"):
            add("loi", "Nhan", r, "muc_ro phải là 'ro' hoặc 'mo_ho' (R8)")
        if r.get("in_scope") not in ("co", "khong"):
            add("loi", "Nhan", r, "in_scope phải là 'co' hoặc 'khong' (R5)")
        if r.get("in_scope") == "co":
            for col in ("tool", "execution_path"):
                if not r.get(col):
                    add("loi", "Nhan", r, f"in_scope=co nhưng thiếu {col}; report không nêu thì ghi 'khong_ro' (R6)")
            if r.get("execution_path") and r["execution_path"] not in ("local", "remote", "khong_ro"):
                add("loi", "Nhan", r, "execution_path phải là local / remote / khong_ro (R6)")
        elif any(r.get(c) for c in ("tool", "execution_path", "source_role", "target_role",
                                     "account_type", "preceded_by", "command_example")):
            add("canh_bao", "Nhan", r, "cột context chỉ điền cho dòng in_scope=co (R6)")

    IOC_TYPES = {"ip", "domain", "url", "hash", "ten_file", "duong_dan", "registry", "named_pipe", "email", "khac"}
    ioc_seen = set()
    for r in data["ioc"]:
        if r["ma_doan"] not in seg_ids:
            add("loi", "IOC", r, "ma_doan không có trong DoanVan")
        if r.get("loai") not in IOC_TYPES:
            add("loi", "IOC", r, f"loai '{r.get('loai')}' không thuộc danh sách của sheet IOC")
        if not r.get("gia_tri"):
            add("loi", "IOC", r, "thiếu gia_tri")
            continue
        k = (r["report"], str(r["gia_tri"]))
        if k in ioc_seen:
            add("loi", "IOC", r, "IOC ghi lặp trong cùng report (R10)")
        ioc_seen.add(k)

    for s in data["segments"]:
        row = {"ma_doan": s.ma_doan, "_row": None}
        if s.loai == "ảnh" and s.anh_hanh_vi not in ("khong", "lap_lai", "bo_sung"):
            add("loi", "DoanVan", row, "đoạn ảnh chưa có cờ anh_hanh_vi (R7)")
    return issues


def progress(data: dict) -> dict[str, dict]:
    """Tiến độ theo report: số đoạn, đã xem, số dòng nhãn, số IOC. Dùng để biết đã sẵn sàng đóng băng chưa."""
    out: dict[str, dict] = {}
    for s in data["segments"]:
        d = out.setdefault(s.report, {"doan": 0, "da_xem": 0, "dong_nhan": 0, "ioc": 0})
        d["doan"] += 1
        d["da_xem"] += int(s.da_xem)
    for r in data["nhan"]:
        out.setdefault(r["report"], {"doan": 0, "da_xem": 0, "dong_nhan": 0, "ioc": 0})["dong_nhan"] += 1
    for r in data["ioc"]:
        out.setdefault(r["report"], {"doan": 0, "da_xem": 0, "dong_nhan": 0, "ioc": 0})["ioc"] += 1
    return out
