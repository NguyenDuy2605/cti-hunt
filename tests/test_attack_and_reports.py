import json

from openpyxl import Workbook

from cti_hunt.attack.catalog import AttackCatalog, tactic_key
from cti_hunt.eval.extraction_metrics import evaluate
from cti_hunt.labels import reports


def ap(sid, tid, name, phases, revoked=False):
    return {"type": "attack-pattern", "id": sid, "name": name, "revoked": revoked,
            "external_references": [{"source_name": "mitre-attack", "external_id": tid}],
            "kill_chain_phases": [{"kill_chain_name": "mitre-attack", "phase_name": p} for p in phases]}


def catalog(tmp_path):
    objs = [
        {"type": "x-mitre-collection", "id": "c", "x_mitre_version": "19.2"},
        {"type": "x-mitre-tactic", "id": "t1", "x_mitre_shortname": "persistence", "name": "Persistence"},
        {"type": "x-mitre-tactic", "id": "t2", "x_mitre_shortname": "lateral-movement", "name": "Lateral Movement"},
        {"type": "x-mitre-tactic", "id": "t3", "x_mitre_shortname": "execution", "name": "Execution"},
        ap("a1", "T1546.003", "Windows Management Instrumentation Event Subscription", ["persistence"]),
        ap("a2", "T1047", "Windows Management Instrumentation", ["execution"]),
        ap("a3", "T1000", "Old", [], revoked=True), ap("a4", "T1001", "Mid", [], revoked=True),
        {"type": "relationship", "id": "r1", "relationship_type": "revoked-by", "source_ref": "a3", "target_ref": "a4"},
        {"type": "relationship", "id": "r2", "relationship_type": "revoked-by", "source_ref": "a4", "target_ref": "a1"},
    ]
    p = tmp_path / "stix.json"
    p.write_text(json.dumps({"type": "bundle", "objects": objs}), encoding="utf-8")
    return AttackCatalog.from_stix(p)


def test_catalog_resolve(tmp_path):
    c = catalog(tmp_path)
    assert c.version == "19.2" and tactic_key("Lateral Movement") == "lateral-movement"
    assert c.resolve("t1047").status == "ok"
    r = c.resolve("T1000")
    assert (r.status, r.final) == ("revoked", "T1546.003")          # thu hồi nối tiếp
    assert c.resolve("T9999").status == "unknown"


def test_extraction_metrics_counts_revoked_before_mapping(tmp_path):
    c = catalog(tmp_path)
    gold = {("SY", "T1546.003"), ("SY", "T1047")}
    res = evaluate(gold, {("SY", "T1000"), ("SY", "T9999")}, c)
    assert res["ma_tho"]["bi_thu_hoi"] == ["SY:T1000->T1546.003"]
    assert res["gop"]["tp"] == 1 and res["gop"]["fp"] == 1 and res["gop"]["fn"] == 1


def make_xlsx(path):
    wb = Workbook()
    ws = wb.active; ws.title = "DoanVan"
    ws.append(["ma_doan", "report", "loai", "van_ban", "da_xem", "anh_hanh_vi", "mo_ta_anh"])
    ws.append(["SY-001", "SY", "tiêu đề", "Raindrop", "x", None, None])
    ws.append(["SY-002", "SY", "đoạn văn", "WMI used to run commands on other hosts", "x", None, None])
    ws.append(["SY-003", "SY", "ảnh", "[ẢNH] ...", None, None, None])
    ws = wb.create_sheet("Nhan")
    ws.append(["ma_doan", "report", "tactic", "technique_id", "technique_name", "cau_trich", "muc_ro", "in_scope",
               "tool", "execution_path", "source_role", "target_role", "account_type", "preceded_by",
               "command_example", "ghi_chu"])
    ws.append(["VD-001", "VD", "Persistence", "T1053.005"])
    ws.append(["SY-002", '=IF(A3="","",LEFT(A3,2))', "Lateral Movement", "T1047",
               "Windows Management Instrumentation", "q", "ro", "co", "WMI", "remote"])
    ws.append(["SY-002", '=IF(A4="","",LEFT(A4,2))', "Persistence", "T1000", "Old", "q", "ro", "co"])
    ws = wb.create_sheet("IOC")
    ws.append(["ma_doan", "report", "loai", "gia_tri", "ghi_chu"])
    ws.append(["SY-002", '=IF(A2="","",LEFT(A2,2))', "domain", "bigtopweb[.]com"])
    wb.save(path)


def test_reports_load_and_check(tmp_path):
    x = tmp_path / "g.xlsx"
    make_xlsx(x)
    data = reports.load(x)
    assert [r["report"] for r in data["nhan"]] == ["SY", "SY"]      # lấy từ ma_doan, bỏ dòng VD
    assert reports.gold_pairs(data["nhan"]) == {("SY", "T1047"), ("SY", "T1000")}
    msgs = [i["noi_dung"] for i in reports.check(data, catalog(tmp_path))]
    assert any("tactic theo mục đích" in m for m in msgs)            # T1047 thuộc Execution
    assert any("thu hồi" in m and "T1546.003" in m for m in msgs)
    assert any("thiếu tool" in m for m in msgs)                      # dòng T1000 in_scope=co
    assert any("anh_hanh_vi" in m for m in msgs)
    assert reports.progress(data)["SY"] == {"doan": 3, "da_xem": 2, "dong_nhan": 2, "ioc": 1}


def test_subtechnique_name_accepts_parent_colon_child(tmp_path):
    """Dạng 'Cha: Con' (như VD-001 và attack.mitre.org) là hợp lệ; tên bản ATT&CK cũ thì vẫn bị báo."""
    catalog(tmp_path)                                       # ghi stix.json mẫu, rồi thêm kỹ thuật cha
    objs = json.loads((tmp_path / "stix.json").read_text(encoding="utf-8"))
    objs["objects"].append(ap("a5", "T1546", "Event Triggered Execution", ["persistence"]))
    (tmp_path / "stix.json").write_text(json.dumps(objs), encoding="utf-8")
    c = AttackCatalog.from_stix(tmp_path / "stix.json")
    seg = reports.Segment("PS-011", "PS", "đoạn văn", "x")
    row = lambda i, name: {"_row": i, "ma_doan": "PS-011", "report": "PS", "tactic": "Persistence",
                           "technique_id": "T1546.003", "technique_name": name, "muc_ro": "ro", "in_scope": "khong"}
    names = ["Windows Management Instrumentation Event Subscription",
             "Event Triggered Execution: Windows Management Instrumentation Event Subscription",
             "Event Triggered Execution: WMI Event Subscription"]
    data = {"segments": [seg], "nhan": [row(i, n) for i, n in enumerate(names)], "ioc": []}
    msgs = [(i["dong"], i["noi_dung"]) for i in reports.check(data, c) if "technique_name" in i["noi_dung"]]
    assert [d for d, _ in msgs] == [2]                      # chỉ tên sai thật bị báo
    assert "Event Triggered Execution: Windows Management Instrumentation Event Subscription" in msgs[0][1]


def test_score_split_removes_codes_written_in_the_report(tmp_path):
    c = catalog(tmp_path)
    from cti_hunt.eval import extraction_metrics as em
    # MS viết sẵn T1047 và mã cũ T1000 (thu hồi -> T1546.003); SY không viết mã nào
    L = em.literal_pairs({"MS": ["Bảng: T1047 WMI; T1000 cũ"], "SY": ["không có mã"]}, c)
    assert L == {("MS", "T1047"), ("MS", "T1000"), ("MS", "T1546.003")}
    gold = {("MS", "T1047"), ("MS", "T1546.003"), ("SY", "T1047")}
    pred = {("MS", "T1047"), ("MS", "T1000"), ("SY", "T1546.003")}      # T1000 quy về T1546.003
    res = em.score(gold, {("SY", "T1047")}, pred, c, L, independent={"MS"})
    assert (res["du"]["gop"]["tp"], res["du"]["gop"]["fp"], res["du"]["gop"]["fn"]) == (2, 1, 1)
    b = res["bo_ma_co_san"]
    assert (b["gop"]["tp"], b["gop"]["fp"], b["gop"]["fn"]) == (0, 1, 1)   # chỉ còn phần SY
    assert b["recall_in_scope"] == {"tp": 0, "mau_so": 1, "thieu": [("SY", "T1047")]}
    assert res["du"]["report_doc_lap"]["tp"] == 2 and res["du"]["report_doc_lap"]["fn"] == 0
    assert "bỏ mã có sẵn" in em.format_score(res)


def _ctx_rows():
    """Ba dòng nhãn dựng tay (không đụng bộ tham chiếu): hai dòng in_scope=co, một dòng ngoài scope."""
    base = dict(source_role="", target_role="", account_type="", preceded_by="", command_example="")
    return [
        {"ma_doan": "SY-002", "report": "SY", "technique_id": "T1047", "in_scope": "co",
         **{**base, "tool": "WMI", "execution_path": "remote", "target_role": "domain controller"}},
        {"ma_doan": "SY-003", "report": "SY", "technique_id": "T1546.003", "in_scope": "co",
         **{**base, "tool": "khong_ro", "execution_path": "local", "preceded_by": "T1000"}},
        {"ma_doan": "PS-001", "report": "PS", "technique_id": "T1047", "in_scope": "khong",
         **{**base, "tool": "", "execution_path": ""}},
    ]


def test_context_normalization_treats_khong_ro_as_empty(tmp_path):
    from cti_hunt.eval import extraction_metrics as em
    assert em.norm_text("  WMI. ") == "wmi"
    assert em.norm_text("khong_ro") == "" == em.norm_text(None)          # D3+D4
    c = catalog(tmp_path)
    assert em.norm_ids("T1000, T1047", c) == ("T1047", "T1546.003")      # mã thu hồi quy đổi trước khi so
    assert em.norm_ids(["T1001"], c) == ("T1546.003",)


def test_context_fields_pair_by_segment_and_technique(tmp_path):
    """Khóa ghép là (đoạn, mã kỹ thuật); dòng ngoài scope chỉ được đếm, không vào P/R (chốt 23/9)."""
    from cti_hunt.eval import extraction_metrics as em
    c = catalog(tmp_path)
    gold = reports.gold_context(_ctx_rows())
    assert set(gold) == {("SY-002", "T1047"), ("SY-003", "T1546.003"), ("PS-001", "T1047")}
    empty = dict(tool="", execution_path="", source_role="", target_role="", account_type="",
                 preceded_by=[], command_example="")
    pred = {
        ("SY-002", "T1047"): {"report": "SY", **{**empty, "tool": "wmi.", "execution_path": "remote",
                                                 "target_role": "the domain controller"}},
        ("SY-003", "T1546.003"): {"report": "SY", **{**empty, "execution_path": "local",
                                                     "preceded_by": ["T1001"]}},
        ("PS-001", "T1047"): {"report": "PS", **{**empty, "tool": "PowerShell"}},
        ("XX-999", "T1047"): {"report": "XX", **empty},                  # đoạn không có trong nhãn
    }
    res = em.score_context(gold, pred, c, literal={("SY", "T1546.003")}, thieu_ma_doan=1)

    g = res["ghep"]
    assert (g["ghep_duoc"], g["dong_nhan"], g["du_doan_thieu_ma_doan"]) == (3, 3, 1)
    assert g["du_doan_khong_ghep_duoc"] == [("XX-999", "T1047")]

    du = res["du"]
    assert du["so_dong"] == 2                                            # chỉ hai dòng in_scope=co
    assert du["truong"]["tool"]["tp"] == 1                               # 'WMI' == 'wmi.' sau chuẩn hóa
    assert du["truong"]["tool"]["ca_hai_trong"] == 1                     # khong_ro vs '' -> không vào mẫu số
    assert du["truong"]["execution_path"]["tp"] == 2
    assert du["truong"]["preceded_by"]["tp"] == 1                        # T1000 và T1001 cùng quy về T1546.003
    tr = du["truong"]["target_role"]
    assert (tr["tp"], tr["fp"], tr["fn"], tr["khop_chua_them"]) == (0, 1, 1, 1)   # chứa nhau -> số phụ
    assert tr["lech"][0]["nhan"] == "domain controller"
    # Nội dung bộ tham chiếu chỉ nằm trong file kết quả, KHÔNG in ra màn hình (docs/METHODOLOGY.md §4.1)
    assert "domain controller" not in em.format_context(res)

    ns = du["ngoai_scope"]                                               # đếm, không precision
    assert (ns["so_dong_ghep"], ns["co_dien_context"], ns["theo_truong"]["tool"]) == (1, 1, 1)
    assert "precision" not in ns

    bo = res["bo_ma_co_san"]                                             # bỏ cặp (SY, T1546.003)
    assert bo["so_dong"] == 1 and bo["truong"]["preceded_by"]["tp"] == 0
    assert "ngoài scope" in em.format_context(res) and "in_scope=co" in em.format_context(res)
